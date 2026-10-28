from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def parse(name, data):
    from app.knowledge.parsing import parse_document

    return parse_document(name, data)


def pdf_bytes(text=None, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def test_text_csv_and_markdown_locations():
    assert parse("a.txt", b"hello\nworld")[0] == ("hello\nworld", "lines 1-2")
    assert parse("a.md", b"# hello")[0][0] == "# hello"
    assert parse("a.csv", b"\xef\xbb\xbfcol,value\na,2")[1] == ("a | 2", "row 2")


def test_real_office_and_pdf():
    document = Document()
    document.add_paragraph("paragraph text")
    out = BytesIO()
    document.save(out)
    assert parse("a.docx", out.getvalue())[0] == ("paragraph text", "paragraph 1")
    workbook = Workbook()
    workbook.active.append(["name", "value"])
    workbook.active.append(["hello", "=1+1"])
    out = BytesIO()
    workbook.save(out)
    result = parse("a.xlsx", out.getvalue())
    assert result[0] == ("name | value", "Sheet row 1")
    assert "=1+1" not in str(result)
    assert parse("a.pdf", pdf_bytes("hello pdf"))[0] == ("hello pdf", "page 1")


@pytest.mark.parametrize(
    "name,data,reason",
    [
        ("../a.txt", b"a", "filename"),
        ("a.html", b"a", "type"),
        ("a.docm", b"a", "type"),
        ("a.txt", b"\xff\xfe", "UTF-8"),
        ("a.pdf", pdf_bytes(), "OCR"),
        ("a.pdf", pdf_bytes("hello", True), "Encrypted"),
        ("a.docx", b"bad", "Office"),
        pytest.param("a.txt", b"a" * (20 * 1024 * 1024 + 1), "20 MiB", id="oversize"),
    ],
)
def test_reject_unsafe_and_unsupported(name, data, reason):
    from app.knowledge.parsing import ParseFailure

    with pytest.raises(ParseFailure, match=reason):
        parse(name, data)


def test_archive_bomb_and_macro_rejected():
    from app.knowledge.parsing import ParseFailure

    for name, data in [("word/vbaProject.bin", b"macro"), ("word/document.xml", b"a" * 1000000)]:
        out = BytesIO()
        with ZipFile(out, "w", ZIP_DEFLATED) as archive:
            archive.writestr(name, data)
        with pytest.raises(ParseFailure):
            parse("a.docx", out.getvalue())


def test_pdf_disables_external_decoders_and_bounds_decompression(monkeypatch):
    import pypdf

    original = pypdf.PdfReader
    observed = []

    def reader(*args, **kwargs):
        observed.append(pypdf.get_configuration())
        return original(*args, **kwargs)

    monkeypatch.setattr(pypdf, "PdfReader", reader)
    with pypdf.apply_configuration(jbig2dec_binary="external-decoder"):
        assert parse("text.pdf", pdf_bytes("text"))
    assert observed[0].jbig2dec_binary is None
    assert observed[0].zlib_maximum_output_length <= 10 * 1024 * 1024


def test_parser_does_not_log_document_content(caplog):
    from app.knowledge.parsing import ParseFailure

    with pytest.raises(ParseFailure):
        parse("private.pdf", b"PRIVATE-CONTENT-invalid-pdf")
    assert "PRIVATE" not in caplog.text


def test_docx_macro_content_type_is_rejected_even_with_renamed_payload():
    from app.knowledge.parsing import ParseFailure

    document = Document()
    document.add_paragraph("ordinary visible text")
    source = BytesIO()
    document.save(source)
    output = BytesIO()
    with ZipFile(source) as old, ZipFile(output, "w") as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/renamed.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            new.writestr(item.filename, data)
        new.writestr("word/renamed.bin", b"macro")
    with pytest.raises(ParseFailure, match="macros"):
        parse("a.docx", output.getvalue())
