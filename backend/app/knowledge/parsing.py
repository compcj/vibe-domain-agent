"""Passive, bounded document extraction. No network, code or formula evaluation."""

import csv
import io
import logging
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

logging.getLogger("pypdf").setLevel(logging.CRITICAL)

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 500000
MAX_PARTS = 10000
MAX_CHUNKS = 1000
ALLOWED = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".xlsx"}


class ParseFailure(ValueError):
    pass


def validate_name(name):
    if (
        not name
        or len(name) > 200
        or any(c in name for c in "/\\:\x00")
        or any(ord(c) < 32 for c in name)
        or name in {".", ".."}
    ):
        raise ParseFailure("Invalid filename")
    suffix = PurePath(name).suffix.lower()
    if suffix not in ALLOWED:
        raise ParseFailure("Unsupported file type")
    return suffix


def guard_office(data):
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(i.file_size for i in entries) > 80 * 1024 * 1024:
                raise ParseFailure("Office archive exceeds extraction limits")
            for item in entries:
                if (
                    item.flag_bits & 1
                    or "vbaproject" in item.filename.lower()
                    or "encryptedpackage" in item.filename.lower()
                ):
                    raise ParseFailure("Encrypted Office files and macros are unsupported")
                if item.file_size > max(1, item.compress_size) * 200:
                    raise ParseFailure("Office archive compression limit exceeded")
            if "[Content_Types].xml" in archive.namelist():
                types = archive.read("[Content_Types].xml").lower()
                if b"vbaproject" in types or b"macroenabled" in types:
                    raise ParseFailure("Office macros are unsupported")
    except BadZipFile:
        raise ParseFailure("Invalid or encrypted Office document") from None


def parse_document(name, data):
    suffix = validate_name(name)
    if len(data) > MAX_BYTES:
        raise ParseFailure("File exceeds 20 MiB limit")
    parts = []
    total = 0

    def append(text, locator):
        nonlocal total
        text = text.strip()
        if not text:
            return
        total += len(text)
        if total > MAX_CHARS or len(parts) >= MAX_PARTS:
            raise ParseFailure("Document exceeds extraction limits")
        parts.append((text, locator))

    try:
        if suffix in {".txt", ".md", ".markdown", ".csv"}:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ParseFailure("Text and CSV files must use UTF-8 encoding") from None
            if "\x00" in text:
                raise ParseFailure("Text and CSV files must use UTF-8 encoding")
            if suffix == ".csv":
                for number, row in enumerate(csv.reader(io.StringIO(text)), 1):
                    if number > MAX_PARTS:
                        raise ParseFailure("Document exceeds row limit")
                    append(" | ".join(row), f"row {number}")
            else:
                append(text, f"lines 1-{len(text.splitlines())}")
        elif suffix == ".pdf":
            parse_pdf(data, append)
            if not parts:
                raise ParseFailure("PDF has no readable text; OCR is not supported")
        else:
            guard_office(data)
            if suffix == ".docx":
                from docx import Document

                document = Document(io.BytesIO(data))
                for number, paragraph in enumerate(document.paragraphs, 1):
                    append(paragraph.text, f"paragraph {number}")
                for number, table in enumerate(document.tables, 1):
                    for row_number, row in enumerate(table.rows, 1):
                        append(
                            " | ".join(cell.text for cell in row.cells),
                            f"table {number} row {row_number}",
                        )
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    io.BytesIO(data), read_only=True, data_only=True, keep_links=False
                )
                try:
                    rows = 0
                    for sheet in workbook:
                        if sheet.max_column and sheet.max_column > 1000:
                            raise ParseFailure("Spreadsheet exceeds column limit")
                        for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                            rows += 1
                            if rows > MAX_PARTS:
                                raise ParseFailure("Spreadsheet exceeds row limit")
                            append(
                                " | ".join("" if cell is None else str(cell) for cell in row),
                                f"{sheet.title} row {number}",
                            )
                finally:
                    workbook.close()
    except ParseFailure:
        raise
    except Exception:
        raise ParseFailure("Document is invalid or cannot be safely parsed") from None
    if not parts:
        raise ParseFailure("Document has no readable text")
    chunks = []
    for text, locator in parts:
        for start in range(0, len(text), 1600):
            chunks.append((text[start : start + 1800], locator))
            if len(chunks) > MAX_CHUNKS:
                raise ParseFailure("Document exceeds chunk limit")
    return chunks


def parse_pdf(data, append):
    from pypdf import PdfReader, apply_configuration

    limit = 10 * 1024 * 1024
    with apply_configuration(
        jbig2dec_binary=None,
        disable_legacy_handling=True,
        maximum_declared_stream_length=limit,
        array_based_stream_maximum_output_length=limit,
        zlib_maximum_output_length=limit,
        lzw_maximum_output_length=limit,
        run_length_maximum_output_length=limit,
        brotli_maximum_output_length=limit,
        image_maximum_buffer_size=limit,
        stream_decoding_work_maximum_length=2 * limit,
        stream_parsing_work_maximum_count=100000,
        page_tree_maximum_depth=30,
        xform_maximum_invocations_per_extraction=500,
    ):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseFailure("Encrypted PDF files are unsupported")
        if len(reader.pages) > 500:
            raise ParseFailure("PDF exceeds page limit")
        for number, page in enumerate(reader.pages, 1):
            content = page.get_contents()
            if content and len(content.get_data()) > limit:
                raise ParseFailure("PDF page exceeds extraction limits")
            append(page.extract_text() or "", f"page {number}")
